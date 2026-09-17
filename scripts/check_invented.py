#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
自创交付物检测 + 素材特征数值残留扫描。

背景（为什么需要）：
    编写响应文件时极易"顺手"发明一些看起来专业、但招标文件并未要求的交付物，
    例如《功能移植确认表》《性能指标实现与验证方案》《建设成效验证报告》。
    这些名字一旦写进投标文件就构成合同义务，会显著加重实施负担、压缩利润，
    甚至因无法交付而被判违约。本脚本把这类"自创交付物"自动筛出来。

同时扫描"素材特征数值残留"：
    参考往期标书时，往往会把素材中的具体指标（质保年限、响应时效、可用率、
    并发规模等）一并带过来，而这些数值通常与本项目招标要求不符，属于隐蔽
    的实质性偏离。本脚本按用户给定的"素材特征值清单"扫描稿中是否残留。

用法：
    python3 check_invented.py --draft "*.md" \
        --sources 技术规格书.txt 招标文件.txt 参考素材A.txt 参考素材B.txt 子系统资料/

    # 扫描素材数值残留
    python3 check_invented.py --draft "*.md" --leak 99.99 30000在线 1年质保

退出码：0=无自创交付物；1=发现疑似自创项（需人工确认后再定稿）
"""
import argparse
import glob
import os
import re
import sys


def load_corpus(paths):
    buf = []
    for p in paths:
        if os.path.isdir(p):
            for root, _, files in os.walk(p):
                for fn in files:
                    fp = os.path.join(root, fn)
                    buf.append(read_any(fp))
        else:
            buf.append(read_any(p))
    return re.sub(r"\s", "", "".join(buf))


def read_any(path):
    """尽力读取文本；.docx/.xlsx 用对应库抽取，其他按纯文本。"""
    low = path.lower()
    try:
        if low.endswith(".docx"):
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    for c in row.cells:
                        parts.append(c.text)
            return "\n".join(parts)
        if low.endswith((".xlsx", ".xlsm")):
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append(" ".join(str(v) for v in row if v is not None))
            return "\n".join(parts)
        if low.endswith(".doc"):
            import subprocess
            return subprocess.run(["textutil", "-convert", "txt", "-stdout", path],
                                  capture_output=True, text=True).stdout
        return open(path, encoding="utf-8", errors="ignore").read()
    except Exception as e:
        print(f"  [跳过] {path}: {e}", file=sys.stderr)
        return ""


def collect_drafts(patterns):
    files = []
    for pat in patterns:
        files.extend(glob.glob(pat))
    return sorted(set(files))


def main():
    ap = argparse.ArgumentParser(description="自创交付物检测 / 素材数值残留扫描")
    ap.add_argument("--draft", nargs="+", required=True,
                    help="待检稿子（支持通配符，如 '*.md'）")
    ap.add_argument("--sources", nargs="*", default=[],
                    help="权威来源：招标文件/技术规格书/参考素材/子系统资料（文件或目录）")
    ap.add_argument("--leak", nargs="*", default=[],
                    help="需扫描的素材特征数值，如 99.99 '30000在线' '1年质保'")
    ap.add_argument("--min-len", type=int, default=2,
                    help="《》名称最小长度（默认2）")
    args = ap.parse_args()

    corpus = load_corpus(args.sources) if args.sources else ""
    files = collect_drafts(args.draft)
    if not files:
        print("未找到待检文件"); return 1

    names, seen = [], set()
    for f in files:
        for m in re.findall(r"《([^》]{1,40})》", open(f, encoding="utf-8").read()):
            key = re.sub(r"\s", "", m)
            if len(key) < args.min_len or key in seen:
                continue
            seen.add(key)
            names.append((key, f))

    invented = [(n, f) for n, f in names if n not in corpus] if corpus else []

    print(f"稿子文件 {len(files)} 个；《》名称 {len(names)} 个")
    if corpus:
        print(f"权威来源语料 {len(corpus)} 字\n")
        print(f"{'名称':<38}{'判定'}")
        print("-" * 58)
        for n, _ in sorted(names, key=lambda x: (x[0] in corpus, x[0])):
            print(f"{n:<38}{'○ 来源已有' if n in corpus else '★ 疑似自创'}")
        print("-" * 58)
        print(f"\n★ 疑似自创交付物 {len(invented)} 个：")
        for n, f in invented:
            print(f"   - {n}   ({os.path.basename(f)})")
        if invented:
            print("\n处置建议：逐一确认。凡招标文件/规格书未要求、且不属于行业惯例交付物的，")
            print("          应去除书名号改为通用表述（如'形成核查记录'），避免形成合同义务。")
    else:
        print("（未提供 --sources，跳过自创判定，仅列出全部《》名称）")
        for n, _ in names:
            print("   -", n)

    if args.leak:
        print(f"\n{'='*58}\n素材特征数值残留扫描\n{'='*58}")
        hit = 0
        for f in files:
            txt = open(f, encoding="utf-8").read()
            for kw in args.leak:
                for m in re.finditer(re.escape(kw), txt):
                    line = txt[:m.start()].count("\n") + 1
                    ctx = txt[max(0, m.start()-40):m.end()+40].replace("\n", " ")
                    print(f"  ⚠ {os.path.basename(f)}:{line}  「{kw}」  …{ctx}…")
                    hit += 1
        print(f"\n共命中 {hit} 处。凡与素材相同但与本项目招标要求不符的数值，必须逐项替换。")

    return 1 if invented else 0


if __name__ == "__main__":
    sys.exit(main())
