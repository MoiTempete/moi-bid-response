#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
直读 markdown 生成 Word —— 绕过 JSON 中间格式，避免超长文档丢内容。

为什么需要这个脚本（替代 generate_docx.py 的 JSON 路径）：
    原 generate_docx.py 要求先把全部正文序列化成 JSON 再喂给脚本。当响应文件
    超过约 1.5 万字 / JSON 超过 50KB 时，序列化与传输极易被截断，表现为
    "文档尾部内容消失"或"部分章节空白"。本脚本直接读取磁盘上的 markdown
    文件，全程无中间序列化，从根本上消除该类丢失。

用法：
    # 1) 生成默认配置模板
    python3 generate_docx_direct.py --init ./doc_config.json

    # 2) 按配置生成（推荐）
    python3 generate_docx_direct.py --config ./doc_config.json

    # 3) 零配置自动模式：把当前目录下 NN-*.md 按序号合成一个文档
    python3 generate_docx_direct.py --auto --title "XX项目技术响应文件"

生成后务必运行 verify_docx.py 做源↔产物双向核对。
"""
import argparse
import glob
import json
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# ----------------------------------------------------------------------------
# 默认排版（可被 doc_config.json 的 format 段整体或逐项覆盖）
# ----------------------------------------------------------------------------
DEFAULT_FORMAT = {
    "body_cn": "仿宋",             # 正文中文字体
    "body_size_pt": 12,            # 小四
    "heading_cn": "黑体",           # 各级标题中文字体
    "heading_size_pt": 14,         # 四号
    "title_cn": "黑体",             # 封面文档标题中文字体
    "title_size_pt": 14,           # 四号
    "en_font": "Times New Roman",  # 全文英文字体
    "line_spacing": 1.5,
    "space_before_pt": 0,
    "space_after_pt": 0,
    "first_line_chars": 2,         # 正文首行缩进字符数（用 firstLineChars，非固定值）
    "table_size_pt": 10.5,
    "table_header_bold": True,
    "add_toc": False,
    "page": {
        "width_cm": 21.0, "height_cm": 29.7,
        "top_cm": 2.5, "bottom_cm": 2.5, "left_cm": 3.2, "right_cm": 3.2,
    },
    # 国标多级自动编号：H2 一、 H3 1. H4 1.1. H5 1.1.1. H6 1.1.1.1.
    "numbering": [
        {"ilvl": 0, "numFmt": "chineseCountingThousand", "lvlText": "%1、"},
        {"ilvl": 1, "numFmt": "decimal", "lvlText": "%2."},
        {"ilvl": 2, "numFmt": "decimal", "lvlText": "%2.%3."},
        {"ilvl": 3, "numFmt": "decimal", "lvlText": "%2.%3.%4."},
        {"ilvl": 4, "numFmt": "decimal", "lvlText": "%2.%3.%4.%5."},
        {"ilvl": 5, "numFmt": "decimal", "lvlText": "%2.%3.%4.%5.%6."},
    ],
    # 章末工作标记（编辑痕迹），生成时丢弃，不计入正文
    "drop_line_patterns": [
        r"^（第[一二三四五六七八九十]+章.*完.*）$",
        r"^（第[一二三四五六七八九十]+章·.*）$",
        r"^【本节完】$",
    ],
}

MD_LINK = re.compile(r"\[([^\]]*)\]\([^)]*\)")
BOLD = re.compile(r"\*\*(.+?)\*\*")


def clean(t: str) -> str:
    """去除 markdown 行内标记，保留文字内容。"""
    t = MD_LINK.sub(r"\1", t)
    t = BOLD.sub(r"\1", t)
    return t.replace("**", "").replace("`", "").strip()


def cn_count(s: str) -> int:
    return len(re.findall(r"[\u4e00-\u9fff]", s))


class DocBuilder:
    def __init__(self, fmt: dict):
        self.f = fmt

    # ---------- 底层工具 ----------
    def _run(self, run, cn_font, size_pt, bold=False):
        run.font.name = self.f["en_font"]
        rPr = run._element.get_or_add_rPr()
        rf = rPr.get_or_add_rFonts()
        rf.set(qn("w:ascii"), self.f["en_font"])
        rf.set(qn("w:hAnsi"), self.f["en_font"])
        rf.set(qn("w:eastAsia"), cn_font)
        run.font.size = Pt(size_pt)
        run.font.bold = bold

    def _para(self, p, first_line_chars=None, align=None, keep_next=False):
        pf = p.paragraph_format
        pf.line_spacing = self.f["line_spacing"]
        pf.space_before = Pt(self.f["space_before_pt"])
        pf.space_after = Pt(self.f["space_after_pt"])
        if align is not None:
            pf.alignment = align
        pf.keep_with_next = keep_next
        pPr = p._element.get_or_add_pPr()
        for old in pPr.findall(qn("w:ind")):
            pPr.remove(old)
        ind = OxmlElement("w:ind")
        if first_line_chars:
            ind.set(qn("w:firstLineChars"), str(int(first_line_chars * 100)))
            ind.set(qn("w:firstLine"),
                      str(int(first_line_chars * self.f["body_size_pt"] * 20)))
        else:
            ind.set(qn("w:firstLineChars"), "0")
            ind.set(qn("w:firstLine"), "0")
        ind.set(qn("w:left"), "0")
        pPr.append(ind)

    # ---------- 文档骨架 ----------
    def setup(self, doc: Document):
        f = self.f
        n = doc.styles["Normal"]
        n.font.name = f["en_font"]
        n.font.size = Pt(f["body_size_pt"])
        n._element.rPr.rFonts.set(qn("w:eastAsia"), f["body_cn"])
        n.paragraph_format.line_spacing = f["line_spacing"]
        n.paragraph_format.space_before = Pt(f["space_before_pt"])
        n.paragraph_format.space_after = Pt(f["space_after_pt"])

        for lvl in range(1, 7):
            name = f"Heading {lvl}"
            try:
                st = doc.styles[name]
            except KeyError:
                continue
            is_doc_title = lvl == 1
            st.font.name = f["en_font"]
            st.font.size = Pt(f["title_size_pt"] if is_doc_title else f["heading_size_pt"])
            st.font.bold = True
            st.font.color.rgb = None
            st._element.rPr.rFonts.set(qn("w:eastAsia"),
                                      f["title_cn"] if is_doc_title else f["heading_cn"])
            pf = st.paragraph_format
            pf.line_spacing = f["line_spacing"]
            pf.space_before = Pt(f["space_before_pt"])
            pf.space_after = Pt(f["space_after_pt"])
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_doc_title else WD_ALIGN_PARAGRAPH.LEFT
            pf.keep_with_next = not is_doc_title
            pPr = st._element.get_or_add_pPr()
            for old in pPr.findall(qn("w:ind")):
                pPr.remove(old)
            ind = OxmlElement("w:ind")
            for k in ("w:firstLineChars", "w:firstLine", "w:left"):
                ind.set(qn(k), "0")
            pPr.append(ind)

        pg = f["page"]
        for s in doc.sections:
            s.page_width, s.page_height = Cm(pg["width_cm"]), Cm(pg["height_cm"])
            s.top_margin = Cm(pg["top_cm"])
            s.bottom_margin = Cm(pg["bottom_cm"])
            s.left_margin = Cm(pg["left_cm"])
            s.right_margin = Cm(pg["right_cm"])
        self.setup_numbering(doc)

    def setup_numbering(self, doc: Document):
        an = OxmlElement("w:abstractNum")
        an.set(qn("w:abstractNumId"), "0")
        an.set(qn("w:restartNumberingAfterBreak"), "0")
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "hybridMultilevel")
        an.append(mlt)
        for cfg in self.f["numbering"]:
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(cfg["ilvl"]))
            lvl.set(qn("w:tplc"), "FFFFFFFF")
            for tag, val in (("w:start", "1"),
                             ("w:numFmt", cfg["numFmt"]),
                             ("w:lvlText", cfg["lvlText"])):
                e = OxmlElement(tag)
                e.set(qn("w:val"), val)
                lvl.append(e)
            pPr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), "0")
            ind.set(qn("w:firstLine"), "0")
            pPr.append(ind)
            lvl.append(pPr)
            an.append(lvl)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), "1")
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), "0")
        num.append(ref)
        ne = doc.part.numbering_part._element
        for t in ("w:abstractNum", "w:num"):
            for old in ne.findall(qn(t)):
                ne.remove(old)
        ne.append(an)
        ne.append(num)

    # ---------- 内容 ----------
    def heading(self, doc, text, level, numbered=True):
        p = doc.add_paragraph(style=f"Heading {min(level, 6)}")
        if numbered and level >= 2:
            pPr = p._element.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            nid = OxmlElement("w:numId")
            nid.set(qn("w:val"), "1")
            numPr.append(nid)
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), str(min(level - 2, len(self.f["numbering"]) - 1)))
            numPr.append(il)
            pPr.append(numPr)
        self._para(p, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        self._run(p.add_run(text), self.f["heading_cn"], self.f["heading_size_pt"], True)
        return p

    def doc_title(self, doc, text):
        p = doc.add_paragraph(style="Heading 1")
        self._para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        parts = text.split("\n")
        r = p.add_run(parts[0])
        self._run(r, self.f["title_cn"], self.f["title_size_pt"], True)
        for extra in parts[1:]:
            r.add_break()
            r.add_text(extra)
        return p

    def body(self, doc, text):
        if not text or not text.strip():
            return None
        p = doc.add_paragraph()
        self._para(p, first_line_chars=self.f["first_line_chars"])
        self._run(p.add_run(text), self.f["body_cn"], self.f["body_size_pt"])
        return p

    def table(self, doc, headers, rows):
        t = doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(headers):
            c = t.rows[0].cells[j]
            p = c.paragraphs[0]
            self._para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
            self._run(p.add_run(h), self.f["heading_cn"], self.f["table_size_pt"],
                      self.f["table_header_bold"])
        for i, row in enumerate(rows, start=1):
            for j in range(len(headers)):
                c = t.rows[i].cells[j]
                p = c.paragraphs[0]
                self._para(p)
                txt = row[j] if j < len(row) else ""
                self._run(p.add_run(txt), self.f["body_cn"], self.f["table_size_pt"])
        return t

    # ---------- markdown 解析 ----------
    def parse(self, path):
        drop = [re.compile(p) for p in self.f.get("drop_line_patterns", [])]
        lines = open(path, encoding="utf-8").read().split("\n")
        blocks, i = [], 0
        while i < len(lines):
            s = lines[i].strip()
            if not s:
                i += 1
                continue
            if any(d.match(s) for d in drop):
                i += 1
                continue
            m = re.match(r"^(#{1,6})\s+(.*)$", s)
            if m:
                blocks.append(("h", len(m.group(1)), clean(m.group(2))))
                i += 1
                continue
            if s.startswith("|") and i + 1 < len(lines) \
                    and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
                headers = [clean(c) for c in s.strip("|").split("|")]
                i += 2
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append([clean(c) for c in lines[i].strip().strip("|").split("|")])
                    i += 1
                blocks.append(("t", headers, rows))
                continue
            blocks.append(("p", clean(s)))
            i += 1
        return blocks

    def build(self, out_path, title, chapters):
        """chapters: [{'file':..., 'title':..., 'numbered':bool}]；title=None 表示无封面。"""
        doc = Document()
        self.setup(doc)
        if title:
            self.doc_title(doc, title)
        for ch in chapters:
            fpath, override = ch["file"], ch.get("title")
            numbered = ch.get("numbered", True)
            for b in self.parse(fpath):
                if b[0] == "h":
                    _, lvl, text = b
                    if lvl == 1:
                        self.heading(doc, override or text, 2, numbered=numbered)
                    else:
                        self.heading(doc, text, lvl + 1, numbered=numbered)
                elif b[0] == "p":
                    self.body(doc, b[1])
                else:
                    self.table(doc, b[1], b[2])
        if self.f.get("add_toc"):
            self._add_toc(doc)
        doc.save(out_path)
        return out_path

    def _add_toc(self, doc):
        p = doc.add_paragraph()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
        p._element.append(fld)


# ----------------------------------------------------------------------------
def load_config(path):
    cfg = json.load(open(path, encoding="utf-8"))
    fmt = dict(DEFAULT_FORMAT)
    fmt.update(cfg.get("format", {}))
    if "page" in cfg.get("format", {}):
        pg = dict(DEFAULT_FORMAT["page"])
        pg.update(cfg["format"]["page"])
        fmt["page"] = pg
    return fmt, cfg.get("documents", [])


def write_default_config(path):
    cfg = {
        "_说明": "format 段可整体删除以使用内置默认；documents 段定义每个输出文档。",
        "format": DEFAULT_FORMAT,
        "documents": [
            {
                "output": "技术响应文件.docx",
                "title": "XX项目\n技术响应文件",
                "chapters": [
                    {"file": "01-项目整体理解与实施思路及推进计划.md",
                     "title": "项目整体理解与实施思路及推进计划"},
                    {"file": "02-项目服务范围、具体工作内容界定.md",
                     "title": "项目服务范围、具体工作内容界定"},
                ],
            }
        ],
    }
    json.dump(cfg, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    return path


def auto_documents(pattern="[0-9][0-9]-*.md"):
    files = sorted(glob.glob(pattern))
    chs = []
    for f in files:
        h1 = ""
        for ln in open(f, encoding="utf-8"):
            m = re.match(r"^#\s+(.*)$", ln.strip())
            if m:
                h1 = clean(m.group(1))
                break
        h1 = re.sub(r"^第[一二三四五六七八九十]+章\s*", "", h1)
        chs.append({"file": f, "title": h1})
    return [{"output": "输出文档.docx", "title": None, "chapters": chs}]


def main():
    ap = argparse.ArgumentParser(description="直读 markdown 生成 Word（无中间格式）")
    ap.add_argument("--config")
    ap.add_argument("--init", metavar="PATH", help="写出默认配置模板后退出")
    ap.add_argument("--auto", action="store_true", help="自动收集 NN-*.md 合成一个文档")
    ap.add_argument("--title", help="--auto 模式下的封面标题")
    ap.add_argument("--format", help="格式配置 JSON 文件（可选，覆盖默认）")
    args = ap.parse_args()

    if args.init:
        print("已写出配置模板:", write_default_config(args.init))
        return 0

    fmt = dict(DEFAULT_FORMAT)
    if args.format:
        fmt.update(json.load(open(args.format, encoding="utf-8")))

    if args.config:
        fmt, docs = load_config(args.config)
    elif args.auto:
        docs = auto_documents()
        docs[0]["title"] = args.title
    else:
        ap.error("需要 --config 或 --auto（或用 --init 生成模板）")

    b = DocBuilder(fmt)
    for d in docs:
        out = d["output"]
        src_cn = sum(cn_count(open(ch["file"], encoding="utf-8").read())
                     for ch in d["chapters"])
        b.build(out, d.get("title"), d["chapters"])
        print(f"[生成] {out}   源文件中文 {src_cn} 字")
    print("\n⚠ 生成后请运行：python3 verify_docx.py --config <同一配置> 做源↔产物核对")
    return 0


if __name__ == "__main__":
    sys.exit(main())
